import hashlib
import json
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

MONEY_COLUMNS = {"sales": 9, "total": 10, "overdue": 12, "our_debt": 14}


class ImportValidationError(ValueError):
    """The source cannot be safely reconciled or interpreted."""


def minor_units(value, location: str = "") -> int:
    if value is None:
        return 0
    if isinstance(value, bool):
        raise ImportValidationError(f"{location}: логическое значение вместо суммы")
    try:
        number = Decimal(str(value).replace("\u00a0", "").replace(" ", "").replace(",", "."))
    except InvalidOperation as exc:
        raise ImportValidationError(f"{location}: некорректная сумма {value!r}") from exc
    if not number.is_finite():
        raise ImportValidationError(f"{location}: сумма должна быть конечной")
    return int((number * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def parse_date(value, location: str = "") -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value).strip(), "%d.%m.%Y").date()
    except ValueError as exc:
        raise ImportValidationError(f"{location}: некорректная дата {value!r}") from exc


def whole_number(value, location: str) -> int | None:
    if value is None or value == "":
        return None
    try:
        number = Decimal(str(value))
    except InvalidOperation as exc:
        raise ImportValidationError(f"{location}: некорректное число дней") from exc
    if not number.is_finite() or number != number.to_integral_value():
        raise ImportValidationError(f"{location}: число дней должно быть целым")
    return int(number)


@dataclass
class ReportNode:
    row: int
    label: str
    indent: int
    parent: "ReportNode | None"
    amounts: dict[str, int]
    created: date | None
    due: date | None
    days_until: int | None
    overdue_days: int
    children: list["ReportNode"] = field(default_factory=list)
    role: str = "unit"

    def ancestors(self) -> list["ReportNode"]:
        result = []
        parent = self.parent
        while parent:
            result.append(parent)
            parent = parent.parent
        return list(reversed(result))


@dataclass
class ParsedReport:
    filename: str
    source_hash: str
    snapshot_date: date
    organization: str
    department_id: str
    department_name: str
    department_source_name: str
    nodes: list[ReportNode]
    leaves: list[ReportNode]
    controls: dict[str, int]


def document_kind(label: str) -> str:
    match = re.match(r"^(.*?)\s+(?:[А-ЯA-Z]{2,8}-\d+|№)", label)
    return match.group(1) if match else label.split(" от ")[0]


def allocation_keys(report: ParsedReport) -> dict[int, str]:
    occurrences: Counter = Counter()
    result = {}
    for node in report.leaves:
        path = [parent.label for parent in node.ancestors()] + [node.label]
        encoded = json.dumps([report.department_id, path], ensure_ascii=False)
        occurrences[encoded] += 1
        # Repeated source allocations remain separate, even with an identical document label.
        result[node.row] = hashlib.sha256(f"{encoded}:{occurrences[encoded]}".encode()).hexdigest()
    return result


def parse_report(path: Path) -> ParsedReport:
    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        if len(workbook.sheetnames) != 1:
            raise ImportValidationError("Ожидался один лист отчёта; нужен явный выбор листа")
        sheet = workbook.active
        reporting_dates = set()
        header_row = None
        total_rows = []
        for row in sheet:
            for cell in row:
                if cell.data_type == "e" and cell.value is not None:
                    raise ImportValidationError(f"{cell.coordinate}: ошибка Excel {cell.value}")
                if isinstance(cell.value, str):
                    match = re.search(r"Дата отчета:\s*(\d{2}\.\d{2}\.\d{4})", cell.value)
                    if match:
                        reporting_dates.add(parse_date(match.group(1), cell.coordinate))
            if str(row[0].value or "").strip() == "Расчетный/плановый документ":
                header_row = row[0].row
            if str(row[0].value or "").strip().casefold() == "итого":
                total_rows.append(row[0].row)
        if len(reporting_dates) != 1 or header_row is None or len(total_rows) != 1:
            raise ImportValidationError("Не найдены однозначные дата, заголовок и общий итог отчёта")
        snapshot = reporting_dates.pop()
        final_row = total_rows[0]
        nodes = []
        stack = []
        for row in range(header_row + 1, final_row):
            label_cell = sheet.cell(row, 1)
            if label_cell.value is None:
                if any(sheet.cell(row, col).value is not None for col in MONEY_COLUMNS.values()):
                    raise ImportValidationError(f"Строка {row}: сумма без наименования")
                continue
            indent = int(label_cell.alignment.indent or 0)
            while stack and stack[-1].indent >= indent:
                stack.pop()
            node = ReportNode(
                row=row, label=str(label_cell.value).strip(), indent=indent,
                parent=stack[-1] if stack else None,
                amounts={key: minor_units(sheet.cell(row, col).value, f"{path.name}:{row}/{col}")
                         for key, col in MONEY_COLUMNS.items()},
                created=parse_date(sheet.cell(row, 5).value, f"E{row}"),
                due=parse_date(sheet.cell(row, 7).value, f"G{row}"),
                days_until=whole_number(sheet.cell(row, 8).value, f"H{row}"),
                overdue_days=whole_number(sheet.cell(row, 13).value, f"M{row}") or 0,
            )
            if node.parent:
                node.parent.children.append(node)
            nodes.append(node)
            stack.append(node)
        roots = [node for node in nodes if node.parent is None]
        if len(roots) != 1:
            raise ImportValidationError("Иерархия должна иметь одну организацию")
        controls = {key: minor_units(sheet.cell(final_row, col).value, f"Итого/{col}")
                    for key, col in MONEY_COLUMNS.items()}
        leaves = []
        for node in reversed(nodes):
            if node.children:
                for key in MONEY_COLUMNS:
                    child_sum = sum(child.amounts[key] for child in node.children)
                    if node.amounts[key] != child_sum:
                        raise ImportValidationError(
                            f"Строка {node.row}: {key} не сходится с дочерними строками; "
                            f"разница {node.amounts[key] - child_sum} в сотых долях"
                        )
                child_roles = {child.role for child in node.children}
                node.role = {"document": "contract", "contract": "counterparty",
                             "counterparty": "manager"}.get(next(iter(child_roles)), "unit") \
                    if len(child_roles) == 1 else "unit"
            else:
                node.role = "document"
                if node.created is None or len(node.ancestors()) < 5:
                    raise ImportValidationError(f"Строка {node.row}: не удалось определить документ")
                if node.created > snapshot:
                    raise ImportValidationError(f"Строка {node.row}: дата возникновения позже даты среза")
                if node.amounts["total"] < 0 or not 0 <= node.amounts["overdue"] <= node.amounts["total"]:
                    raise ImportValidationError(f"Строка {node.row}: некорректный долг или просрочка")
                if node.amounts["total"] > 0 and node.due is None:
                    raise ImportValidationError(f"Строка {node.row}: у долга отсутствует плановая дата")
                if node.days_until is not None:
                    if node.due is None or node.days_until != (node.due - snapshot).days:
                        raise ImportValidationError(f"Строка {node.row}: дата и число дней не совпадают")
                if node.amounts["total"] > 0 and node.due and node.due > snapshot \
                        and node.days_until is None:
                    raise ImportValidationError(f"Строка {node.row}: отсутствуют дни до погашения")
                if node.amounts["overdue"] > 0 and node.overdue_days <= 0:
                    raise ImportValidationError(f"Строка {node.row}: отсутствуют дни просрочки")
                leaves.append(node)
        leaves.reverse()
        for key in MONEY_COLUMNS:
            if sum(node.amounts[key] for node in leaves) != controls[key] \
                    or roots[0].amounts[key] != controls[key]:
                raise ImportValidationError(f"Общий итог {key} не сходится с документами")
        department_labels = {node.ancestors()[2].label for node in leaves}
        if len(department_labels) != 1:
            raise ImportValidationError("Ожидался один отдел в файле")
        department_label = department_labels.pop()
        if "комплектац" in department_label.casefold():
            department_id, department_name = "completion", "Комплектация"
        elif "дистрибьюц" in department_label.casefold():
            department_id, department_name = "distribution", "Дистрибьюция"
        else:
            raise ImportValidationError(f"Неизвестный отдел: {department_label}")
        for node in leaves:
            if node.parent.role != "contract" or node.parent.parent.role != "counterparty":
                raise ImportValidationError(f"Строка {node.row}: неоднозначная связь клиента и договора")
        roots[0].role = "organization"
        return ParsedReport(path.name, hashlib.sha256(path.read_bytes()).hexdigest(), snapshot,
                            roots[0].label, department_id, department_name, department_label,
                            nodes, leaves, controls)
    finally:
        workbook.close()
